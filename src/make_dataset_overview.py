"""
make_dataset_overview.py - every dataset in the portfolio, in one Excel workbook to browse.

Output: data/VMG_datasets_overview.xlsx (one sheet per dataset, plus a contents sheet).
Run after make_exports.py, fetch_external.py and reference_kpis.py:
  python src/make_dataset_overview.py
"""

from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

import reference_kpis as rk

ROOT = Path(__file__).resolve().parents[1]
DATA = ROOT / "data"
OUT = DATA / "VMG_datasets_overview.xlsx"

HEADER_FILL = PatternFill("solid", fgColor="F1F0EC")
ACCENT = "2A78D6"


def main() -> None:
    raw = rk.load_exports()
    fx = pd.read_csv(DATA / "external" / "fx_rates_monthly.csv")
    reasons = pd.read_csv(DATA / "reference" / "termination_reasons.csv")
    snap = rk.clean(raw, fx)
    mv = rk.movements(snap, reasons)
    kpis = pd.read_csv(DATA / "reference" / "reference_kpis.csv")
    latest = sorted(raw["File"].unique())[-1]

    company = kpis[kpis["Scope"] == "Company"].pivot(index="Month", columns="KPI", values="Value").reset_index()
    sheets = {
        "Export Sep 2026": (raw[raw["File"] == latest].drop(columns=["As Of"]),
                            "One monthly HRIS export, exactly as delivered (SuccessFactors-style columns, synthetic people).", "Synthetic"),
        "All exports": (raw.drop(columns=["As Of"]),
                        "All 25 monthly exports stacked (Sep 2024 - Sep 2026), raw, errors included.", "Synthetic"),
        "Planted issues": (pd.read_csv(DATA / "reference" / "planted_issues.csv", dtype=str),
                           "Answer key: the 19 data errors planted on purpose. The DQ rules must find exactly these.", "Synthetic"),
        "Clean snapshot": (snap.drop(columns=["File", "Month"]).assign(**{"As Of": snap["As Of"].dt.date,
                                                                          "Hire Date": snap["Hire Date"].dt.date,
                                                                          "Event Date": snap["Event Date"].dt.date}),
                           "After cleaning: one row per person per month-end, names removed, EUR salaries added.", "Derived"),
        "Movements": (mv, "Hires, leavers and transfers, found by comparing consecutive month-ends.", "Derived"),
        "KPIs by month": (company, "Company KPIs per month-end (the reference numbers Power BI is checked against).", "Derived"),
        "Eurostat pay gap": (pd.read_csv(DATA / "external" / "eurostat_gender_pay_gap.csv"),
                             "REAL: unadjusted gender pay gap %, Eurostat earn_gr_gpgr2 (p = provisional, d = definition differs, b = break).", "Real"),
        "ECB FX rates": (fx, "REAL: ECB monthly average exchange rates (units per EUR).", "Real"),
        "Legal entities": (pd.read_csv(DATA / "reference" / "legal_entities.csv"), "Lookup: legal entity -> country, currency.", "Reference"),
        "Leaver reasons": (reasons, "Lookup: termination reason -> voluntary / involuntary / retirement.", "Reference"),
        "Ticket 4711 list": (pd.read_csv(DATA / "ticket" / "hrbp_operations_list_2026-10-02.csv", dtype=str),
                             "The HR Business Partner's own Operations list (55 names) for the ticket case.", "Synthetic"),
    }

    contents = pd.DataFrame([{"Sheet": name, "Type": kind, "Rows": len(df), "Columns": df.shape[1], "What it is": desc}
                             for name, (df, desc, kind) in sheets.items()])

    with pd.ExcelWriter(OUT, engine="openpyxl") as xl:
        contents.to_excel(xl, sheet_name="Contents", index=False, startrow=3)
        for name, (df, _, _) in sheets.items():
            df.to_excel(xl, sheet_name=name, index=False)

        ws = xl.sheets["Contents"]
        ws["A1"] = "Vltava Market Group - People Analytics portfolio datasets"
        ws["A1"].font = Font(size=14, bold=True)
        ws["A2"] = "Fictional company. 'Synthetic' = generated people; 'Real' = public Eurostat / ECB data."
        ws["A2"].font = Font(color="52514E", italic=True)
        for row in range(5, 5 + len(contents)):
            cell = ws.cell(row=row, column=1)
            cell.hyperlink = f"#'{cell.value}'!A1"
            cell.font = Font(color=ACCENT, underline="single")

        for name, sheet in xl.sheets.items():
            header_row = 4 if name == "Contents" else 1
            for cell in sheet[header_row]:
                cell.font = Font(bold=True)
                cell.fill = HEADER_FILL
                cell.alignment = Alignment(vertical="center")
            sheet.freeze_panes = sheet.cell(row=header_row + 1, column=1)
            if sheet.max_row > header_row:
                sheet.auto_filter.ref = f"A{header_row}:{get_column_letter(sheet.max_column)}{sheet.max_row}"
            for col in range(1, sheet.max_column + 1):
                values = [sheet.cell(row=r, column=col).value for r in range(header_row, min(sheet.max_row, 300) + 1)]
                width = max(len(str(v)) for v in values if v is not None) if any(v is not None for v in values) else 8
                sheet.column_dimensions[get_column_letter(col)].width = min(max(width + 2, 8), 60 if name == "Contents" else 34)
    print(f"Wrote {OUT.relative_to(ROOT)}")
    print(contents[["Sheet", "Type", "Rows", "Columns"]].to_string(index=False))


if __name__ == "__main__":
    main()
