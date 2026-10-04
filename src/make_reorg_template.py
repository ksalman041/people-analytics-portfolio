"""
make_reorg_template.py - D3: the Excel form HR Business Partners fill in for a reorganisation.

It catches the blocking issues while they type: unknown target unit, wrong cost center,
inactive unit head, wrong effective date, unmapped units and a headcount that doesn't
reconcile. It is pre-filled with the first (v1) submission, so the flags are visible.

Output: excel/reorg_mapping_template.xlsx
Run:    python src/make_reorg_template.py
"""

from __future__ import annotations

from datetime import date
from pathlib import Path

import pandas as pd
from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

ROOT = Path(__file__).resolve().parents[1]
DIR = ROOT / "data" / "reorg"
OUT = ROOT / "excel" / "reorg_mapping_template.xlsx"

FONT = "Segoe UI"
INK, INK_2, ACCENT, GOOD, BAD = "0B0B0B", "52514E", "2A78D6", "006300", "D03B3B"
INPUT_FILL = PatternFill("solid", fgColor="EEF4FC")  # light blue = you type here
CALC_FILL = PatternFill("solid", fgColor="F4F4F1")  # grey = calculated, don't type
BAD_FILL = PatternFill("solid", fgColor="FBE9E9")
HEADER_BORDER = Border(bottom=Side(style="medium", color=ACCENT))
FIRST, LAST = 5, 16  # mapping input rows
PROGRAMME_DATE = date(2026, 7, 1)


def font(size=10, bold=False, color=INK, italic=False):
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def header(ws, row, labels, widths):
    for col, (label, width) in enumerate(zip(labels, widths), start=1):
        c = ws.cell(row=row, column=col, value=label)
        c.font = font(bold=True)
        c.border = HEADER_BORDER
        c.alignment = Alignment(vertical="bottom", wrap_text=True)
        ws.column_dimensions[c.column_letter].width = width
    ws.row_dimensions[row].height = 30


def title(ws, text, sub):
    ws["A1"] = text
    ws["A1"].font = font(16, bold=True)
    ws["A2"] = sub
    ws["A2"].font = font(10, color=INK_2, italic=True)
    ws.sheet_view.showGridLines = False


def main() -> None:
    units = pd.read_csv(DIR / "units_current.csv", dtype=str, keep_default_na=False)
    target = pd.read_csv(DIR / "units_target.csv", dtype=str, keep_default_na=False)
    v1 = pd.read_csv(DIR / "mapping_v1_hrbp.csv", dtype=str, keep_default_na=False)
    positions = pd.read_csv(DIR / "positions_before.csv", dtype=str, keep_default_na=False)
    people = pd.read_csv(DIR / "people_all_before.csv", dtype=str, keep_default_na=False)

    wb = Workbook()

    # --- Lists (reference data the drop-downs and checks use) ----------------
    ls = wb.active
    ls.title = "Lists"
    title(ls, "Reference lists", "Maintained by People Analytics. Drop-downs and checks read from here.")
    header(ls, 4, ["Current unit", "Division", "Cost center", "Head", "", "Target unit", "Target cost center",
                   "Target division", "", "Programme date"], [24, 14, 12, 10, 3, 26, 14, 14, 3, 15])
    for i, r in enumerate(units.itertuples(index=False), start=5):
        for col, v in enumerate([r.Department, r.Division, r[2], int(r.Head)], start=1):
            ls.cell(row=i, column=col, value=v).font = font()
    for i, r in enumerate(target.itertuples(index=False), start=5):
        for col, v in enumerate([r.Department, r[2], r.Division], start=6):
            ls.cell(row=i, column=col, value=v).font = font()
    ls["J5"] = PROGRAMME_DATE
    ls["J5"].number_format = "dd mmm yyyy"
    ls["J5"].font = font()
    n_units, n_target = 4 + len(units), 4 + len(target)
    cur = f"Lists!$A$5:$A${n_units}"
    tgt, tgt_cc = f"Lists!$F$5:$F${n_target}", f"Lists!$G$5:$G${n_target}"

    # --- People (active employees on 30 June, IDs only: no names needed) ------
    pp = wb.create_sheet("People")
    title(pp, "Active employees, 30 June 2026", "Pseudonymous IDs only (data minimisation). Used to check unit heads.")
    header(pp, 4, ["Person ID", "Job title", "Manager ID", "Department (if in scope)"], [12, 36, 12, 26])
    dept = dict(zip(positions["Person ID External"], positions["Department"]))
    for i, r in enumerate(people.itertuples(index=False), start=5):
        pid = r[0]
        pp.cell(row=i, column=1, value=int(pid)).font = font()
        pp.cell(row=i, column=2, value=r[2]).font = font()
        pp.cell(row=i, column=3, value=int(r[1]) if r[1] else None).font = font()
        pp.cell(row=i, column=4, value=dept.get(pid, "")).font = font()
    people_ids = f"People!$A$5:$A${4 + len(people)}"
    people_dept = f"People!$D$5:$D${4 + len(people)}"

    # --- Mapping (the form) ---------------------------------------------------
    ws = wb.create_sheet("Mapping", 0)
    title(ws, "Reorganisation mapping: Technology, Q3 2026",
          "Blue cells: you type or pick. Grey cells: checks, don't type. Every row must say 'Ready' before go-live.")
    header(ws, 4, ["Old unit", "Old cost center", "New division", "New unit", "New cost center", "New unit head (ID)",
                   "Head reports to (ID)", "Effective date", "New unit in target?", "Cost center OK?", "Head active?",
                   "Date OK?", "Status"], [24, 12, 13, 26, 13, 12, 12, 13, 18, 18, 14, 22, 16])
    rows = list(v1.itertuples(index=False))
    for r in range(FIRST, LAST + 1):
        data = rows[r - FIRST] if r - FIRST < len(rows) else None
        values = {}
        if data is not None:
            values = {1: data[0], 3: data[2], 4: data[3], 5: data[4], 6: int(data[5]), 7: int(data[6]),
                      8: date.fromisoformat(data[7])}
        for col in (1, 3, 4, 5, 6, 7, 8):
            c = ws.cell(row=r, column=col, value=values.get(col))
            c.fill, c.font = INPUT_FILL, font()
        ws.cell(row=r, column=8).number_format = "dd mmm yyyy"
        formulas = {
            2: f'=IF(A{r}="","",IFERROR(_xlfn.XLOOKUP(A{r},{cur},Lists!$C$5:$C${n_units}),"? unknown unit"))',
            9: f'=IF(D{r}="","",IF(COUNTIF({tgt},D{r})>0,"✓","✗ not in target"))',
            10: f'=IF(E{r}="","",IF(COUNTIF({tgt},D{r})=0,"–",IF(_xlfn.XLOOKUP(D{r},{tgt},{tgt_cc})=E{r},"✓",'
                f'"✗ use "&_xlfn.XLOOKUP(D{r},{tgt},{tgt_cc}))))',
            11: f'=IF(F{r}="","",IF(COUNTIF({people_ids},F{r})>0,"✓","✗ not found"))',
            12: f'=IF(H{r}="","",IF(H{r}=Lists!$J$5,"✓","✗ should be "&TEXT(Lists!$J$5,"dd mmm yyyy")))',
            13: f'=IF(A{r}="","",IF(COUNTIF(I{r}:L{r},"✗*")=0,"Ready","Fix "&COUNTIF(I{r}:L{r},"✗*")&" issue(s)"))',
        }
        for col, f in formulas.items():
            c = ws.cell(row=r, column=col, value=f)
            c.fill, c.font = CALC_FILL, font()
    rng = f"I{FIRST}:M{LAST}"
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(I{FIRST},1)="✗"'], font=Font(name=FONT, color=BAD, bold=True), fill=BAD_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'LEFT(I{FIRST},3)="Fix"'], font=Font(name=FONT, color=BAD, bold=True), fill=BAD_FILL))
    ws.conditional_formatting.add(rng, FormulaRule(formula=[f'OR(I{FIRST}="✓",I{FIRST}="Ready")'], font=Font(name=FONT, color=GOOD, bold=True)))
    dv_unit = DataValidation(type="list", formula1=f"={cur}", allow_blank=True, errorTitle="Unknown unit",
                             error="Pick a current unit from the list.")
    dv_target = DataValidation(type="list", formula1=f"={tgt}", allow_blank=True, errorStyle="warning",
                               errorTitle="Not in the target structure", error="This unit is not in the approved target structure.")
    dv_cc = DataValidation(type="list", formula1=f"={tgt_cc}", allow_blank=True, errorStyle="warning")
    dv_div = DataValidation(type="list", formula1='"Technology,Markets,Operations,Corporate Functions"', allow_blank=True)
    dv_date = DataValidation(type="date", operator="greaterThan", formula1="DATE(2026,1,1)", allow_blank=True)
    for dv, col in ((dv_unit, "A"), (dv_div, "C"), (dv_target, "D"), (dv_cc, "E"), (dv_date, "H")):
        ws.add_data_validation(dv)
        dv.add(f"{col}{FIRST}:{col}{LAST}")
    ws.freeze_panes = "B5"

    # --- Coverage: every current unit mapped, headcount reconciles -------------
    cv = wb.create_sheet("Coverage", 1)
    title(cv, "Coverage check", "Every unit in scope must be mapped, and headcount must be the same before and after.")
    header(cv, 4, ["Current unit", "Cost center", "Positions", "Mapped?"], [24, 12, 11, 18])
    for i in range(len(units)):
        r = 5 + i
        cv.cell(row=r, column=1, value=f"=Lists!A{5 + i}")
        cv.cell(row=r, column=2, value=f"=Lists!C{5 + i}")
        cv.cell(row=r, column=3, value=f"=COUNTIF({people_dept},A{r})")
        cv.cell(row=r, column=4, value=f'=IF(COUNTIF(Mapping!$A${FIRST}:$A${LAST},A{r})>0,"✓ mapped","✗ NOT MAPPED")')
        for col in range(1, 5):
            cv.cell(row=r, column=col).fill, cv.cell(row=r, column=col).font = CALC_FILL, font()
    end = 4 + len(units)
    summary = [("Positions before", f"=SUM(C5:C{end})"),
               ("Positions mapped", f'=SUMIF(D5:D{end},"✓ mapped",C5:C{end})'),
               ("Headcount reconciles?", f'=IF(B{end + 2}=B{end + 3},"✓ yes","✗ "&B{end + 2}-B{end + 3}&" positions unmapped")')]
    for j, (label, f) in enumerate(summary):
        cv.cell(row=end + 2 + j, column=1, value=label).font = font(bold=True)
        c = cv.cell(row=end + 2 + j, column=2, value=f)
        c.fill, c.font = CALC_FILL, font(bold=True)
    cv.conditional_formatting.add(f"A5:D{end + 4}", FormulaRule(formula=['LEFT(A5,1)="✗"'], font=Font(name=FONT, color=BAD, bold=True), fill=BAD_FILL))
    cv.conditional_formatting.add(f"A5:D{end + 4}", FormulaRule(formula=['LEFT(A5,1)="✓"'], font=Font(name=FONT, color=GOOD, bold=True)))

    # --- Read me ---------------------------------------------------------------
    rd = wb.create_sheet("Read me", 0)
    title(rd, "How to use this template", "Vltava Market Group · People Analytics · synthetic data")
    rd.column_dimensions["A"].width = 110
    lines = [
        ("1. On 'Mapping', add one row per current unit: where does it go, which cost center, who leads it.", False),
        ("2. Blue cells are yours. Grey cells check your entries as you type: every row must say 'Ready'.", False),
        ("3. 'Coverage' shows units you haven't mapped yet and whether headcount still adds up.", False),
        ("4. Send the file to People Analytics. src/reorg_validate.py re-runs these checks and adds spans of control,", False),
        ("   management layers and a comparison with the HRIS after go-live.", False),
        ("", False),
        ("What the checks catch", True),
        ("✗ not in target      the new unit isn't in the approved target structure (often a typo)", False),
        ("✗ use CC-…           the cost center doesn't belong to that unit", False),
        ("✗ not found          the unit head isn't an active employee", False),
        ("✗ should be …        the effective date differs from the programme date", False),
        ("✗ NOT MAPPED         a current unit has no row; its people would drop out of reporting", False),
        ("", False),
        ("This copy is pre-filled with the first HRBP submission, so you can see the flags in action.", True),
    ]
    for i, (text, bold) in enumerate(lines, start=4):
        rd.cell(row=i, column=1, value=text).font = font(10, bold=bold)

    OUT.parent.mkdir(parents=True, exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT.relative_to(ROOT)}")


if __name__ == "__main__":
    main()
